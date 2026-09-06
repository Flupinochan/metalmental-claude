"""Builders for the small real documents the extraction tests need."""

from pathlib import Path

import pytest


def write_docx(path: Path, heading: str, paragraphs: list[str]) -> Path:
    from docx import Document

    document = Document()
    document.add_heading(heading, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))
    return path


def write_pptx(path: Path, title: str, body: str) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title_placeholder = slide.shapes.title
    assert title_placeholder is not None
    title_placeholder.text = title
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = body
    presentation.save(str(path))
    return path


def write_pdf(path: Path, pages: list[str]) -> Path:
    """Emit a minimal one-font PDF so tests do not depend on a checked-in binary."""
    objects, kids = [], []
    for index, text in enumerate(pages):
        content_id = 4 + index * 2
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        objects.append(
            (content_id, b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
        )
        objects.append(
            (
                content_id + 1,
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
                b"/Resources << /Font << /F1 3 0 R >> >> >>" % content_id,
            )
        )
        kids.append(content_id + 1)

    header = [
        (1, b"<< /Type /Catalog /Pages 2 0 R >>"),
        (
            2,
            b"<< /Type /Pages /Kids [%s] /Count %d >>"
            % (b" ".join(b"%d 0 R" % k for k in kids), len(kids)),
        ),
        (3, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"),
    ]
    everything = sorted(header + objects)
    out, offsets = bytearray(b"%PDF-1.4\n"), {}
    for number, body in everything:
        offsets[number] = len(out)
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(everything) + 1)
    for number, _ in everything:
        out += b"%010d 00000 n \n" % offsets[number]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(everything) + 1,
        xref,
    )
    path.write_bytes(bytes(out))
    return path


@pytest.fixture
def extracted_root(tmp_path):
    root = tmp_path / "extracted"
    root.mkdir()
    return root
